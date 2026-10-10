"""
Classification reports (Grzegorz Gutek, 2026-10-07) and the manager history.

The first half needs nothing: configuration, paths, the workbook layout and
the product lookup. The second half needs a real PostgreSQL, like
test_update_uploads.py:

    TEST_DATABASE_URL=postgresql+asyncpg://user:pass@host:5432/rfm_test pytest tests/test_reports.py
"""

import io
import json
import os
import subprocess
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from api.services import report_service
from api.services.report_service import (
    ReportConfigError,
    ReportEntry,
    build_workbook,
    created_by_rfm,
    parse_report_config,
    report_location,
)

WARSAW = ZoneInfo("Europe/Warsaw")


def config(**values):
    return parse_report_config(values)


def entry(op_id, name, pushed_at, product=None, username="mizn"):
    return ReportEntry(op_id, name, pushed_at, username, 3, product or {})


# ---------------------------------------------------------------------------
# Configuration and paths
# ---------------------------------------------------------------------------

def test_defaults_follow_the_request():
    cfg = config()
    assert cfg.auto_enabled is False
    assert [r["name"] for r in cfg.recipients] == ["Natalia", "Lena", "Ewa"]
    assert [c["header"] for c in cfg.columns] == [
        "Nazwa", "Projektant", "Płeć", "Data premiery", "Data ostatniej wysyłki",
    ]
    assert (cfg.day_row_color, cfg.late_row_color) == ("AFD095", "FFA6A6")
    assert str(cfg.tz) == "Europe/Warsaw"
    # Managers see the settings, never the share password or the API key
    assert "smb_password" not in cfg.public() and "product_api_key" not in cfg.public()


def test_october_2026_lands_where_the_email_says():
    cfg = config()
    folder, name = report_location(cfg, date(2026, 10, 1), {"name": "Natalia", "username": "mizn"})
    assert folder == r"\\radius1\Users\Iza.Horna\WAŻNE FOLDERY BEATA GRZESIEK\!RAPORTY KLASYFIKACJA\2026\10 PAŹDZIERNIK"
    assert name == "Raport_klasyfikacji_Natalia_2026_10.xlsx"


def test_placeholder_values_cannot_add_folders():
    cfg = config(reports_output_dir="//srv/share/{name}", reports_file_name="{username}")
    folder, name = report_location(cfg, date(2026, 1, 1), {"name": "a/b\\c", "username": "x:y"})
    assert folder == r"\\srv\share\a_b_c"
    assert name == "x_y.xlsx"


@pytest.mark.parametrize("key, value", [
    ("reports_day_row_color", "green"),
    ("reports_run_time", "25:00"),
    ("reports_timezone", "Mars/Olympus"),
    ("reports_recipients", "{not json"),
    ("reports_recipients", '[{"name": "no username"}]'),
    ("reports_columns", '[{"header": "X", "field": "price"}]'),
    ("reports_columns", "[]"),
    ("reports_month_names", "JAN,FEB"),
])
def test_unusable_settings_name_themselves(key, value):
    with pytest.raises(ReportConfigError, match=key):
        config(**{key: value})


def test_output_dir_must_be_a_share():
    with pytest.raises(ReportConfigError, match="UNC"):
        report_location(config(reports_output_dir="C:\\reports"), date(2026, 10, 1), {"name": "a", "username": "a"})


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def test_days_are_merged_green_rows_and_late_products_are_red():
    cfg = config()
    utc = timezone.utc
    entries = [
        # 1 Oct: one on time, one pushed after its release date
        entry(1, "TORBA A", datetime(2026, 10, 1, 8, 0, tzinfo=utc),
              {"designer": "PRADA", "gender": "K", "release_date": "2026-10-01", "last_send_date": "2026-09-30"}),
        entry(2, "TORBA B", datetime(2026, 10, 1, 9, 0, tzinfo=utc),
              {"designer": "GUCCI", "gender": "M", "release_date": "2026-09-15 00:00:00.000"}),
        # 23:30 UTC on 1 Oct is 2 Oct in Warsaw
        entry(3, "BUTY C", datetime(2026, 10, 1, 23, 30, tzinfo=utc), {"release_date": "2026-12-01"}),
    ]
    data = build_workbook(entries, cfg)
    sheet = load_workbook(io.BytesIO(data)).active

    assert [c.value for c in sheet[1]] == [c["header"] for c in cfg.columns]
    assert sorted(str(r) for r in sheet.merged_cells.ranges) == ["A2:E2", "A5:E5"]
    assert sheet["A2"].value.date() == date(2026, 10, 1)
    assert sheet["A5"].value.date() == date(2026, 10, 2)
    assert sheet["A2"].fill.fgColor.rgb.endswith("AFD095")

    assert [sheet.cell(row=r, column=1).value for r in (3, 4, 6)] == ["TORBA A", "TORBA B", "BUTY C"]
    assert sheet["B3"].value == "PRADA"
    assert sheet["D3"].value.date() == date(2026, 10, 1) and sheet["D3"].number_format == "DD.MM.YYYY"
    assert sheet["A3"].fill.fill_type is None          # released that day: not late
    assert sheet["A4"].fill.fgColor.rgb.endswith("FFA6A6")  # released 15 Sep, pushed 1 Oct
    assert sheet["E4"].fill.fgColor.rgb.endswith("FFA6A6")  # the whole row
    assert sheet["A6"].fill.fill_type is None
    assert created_by_rfm(data) is True


def test_no_product_data_leaves_cells_blank_and_nothing_late():
    data = build_workbook([entry(1, "TORBA A", datetime(2026, 10, 1, 8, tzinfo=timezone.utc))], config())
    sheet = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in sheet[3]] == ["TORBA A", None, None, None, None]
    assert sheet["A3"].fill.fill_type is None


def test_a_file_someone_else_made_is_recognised():
    buffer = io.BytesIO()
    Workbook().save(buffer)
    assert created_by_rfm(buffer.getvalue()) is False
    assert created_by_rfm(b"not a zip") is False


# ---------------------------------------------------------------------------
# Product details service
# ---------------------------------------------------------------------------

@pytest.fixture
def polka(monkeypatch):
    """RFM_ProductDetails answering from a dict; records the names asked for."""
    asked = []
    products = {
        "TORBA A": {"success": True, "found": True, "product": {"designer": "PRADA"}},
        "TORBA B": {"success": True, "found": False, "product": None},
    }

    def handler(request):
        name = request.url.params["CatalogName"]
        asked.append(name)
        assert request.url.params["ApiKey"] == "key"
        assert "accept-charset" not in request.headers
        body = products.get(name, {"success": False, "error": "boom"})
        return httpx.Response(200, content=json.dumps(body).encode("cp1250"))

    real = httpx.AsyncClient
    monkeypatch.setattr(
        report_service.httpx, "AsyncClient",
        lambda **kwargs: real(transport=httpx.MockTransport(handler), **kwargs),
    )
    return asked


async def test_products_are_looked_up_once_per_catalog(polka):
    cfg = config(reports_product_url="http://polka/RFM_ProductDetails", reports_product_api_key="key")
    products, warning = await report_service.fetch_products(["TORBA A", "TORBA A", "TORBA B", "X"], cfg)

    assert sorted(polka) == ["TORBA A", "TORBA B", "X"]
    assert products == {"TORBA A": {"designer": "PRADA"}}
    assert warning.startswith("No product data for 2 of 3 catalog(s)")
    assert "TORBA B: not found" in warning and "X: boom" in warning


async def test_without_the_service_the_report_says_so(polka):
    products, warning = await report_service.fetch_products(["TORBA A"], config())
    assert products == {} and polka == []
    assert "not configured" in warning


# ---------------------------------------------------------------------------
# PostgreSQL: which pushes a report lists, the nightly job, manager history
# ---------------------------------------------------------------------------

TEST_DATABASE_URL = os.getenv("TEST_DATABASE_URL")
needs_db = pytest.mark.skipif(not TEST_DATABASE_URL, reason="TEST_DATABASE_URL not set")


@pytest.fixture(scope="module")
def migrated_database():
    subprocess.run(
        [sys.executable, "-m", "alembic", "upgrade", "head"],
        cwd=Path(__file__).resolve().parent.parent,
        env={**os.environ, "DATABASE_URL": TEST_DATABASE_URL},
        check=True,
        capture_output=True,
    )


@pytest.fixture
async def db(migrated_database, monkeypatch):
    from sqlalchemy import text
    from database import DatabaseManager

    monkeypatch.setenv("DATABASE_URL", TEST_DATABASE_URL)
    await DatabaseManager.close()
    DatabaseManager.initialize(pool_size=5, max_overflow=5)
    async with DatabaseManager.session() as s:
        await s.execute(text("SET LOCAL lock_timeout = '10s'"))
        await s.execute(text(
            "TRUNCATE users, operations, report_runs, pim_events, remote_sync_checks, audit_logs "
            "RESTART IDENTITY CASCADE"
        ))
        # The seeded settings (reports off); other suites empty the config table
        await s.execute(text("DELETE FROM config WHERE key LIKE 'reports_%'"))
        for key, value in report_service.DEFAULTS.items():
            await s.execute(
                text("INSERT INTO config (key, value, type) VALUES (:key, :value, 'STRING')"),
                {"key": key, "value": value},
            )
    yield DatabaseManager
    await DatabaseManager.close()


async def add_user(s, username, role="USER"):
    from models import User, UserRole
    user = User(username=username, password_hash="x", role=UserRole(role))
    s.add(user)
    await s.flush()
    return user


async def add_push(s, user, name, completed_at, status="COMPLETED", files=3):
    from models import Operation, OperationStatus, OperationType
    op = Operation(
        user_id=user.id, type=OperationType.PUSH, status=OperationStatus(status),
        source_path=f"A:/DO KATALOGU/{user.username}/{name}", dest_path=f"B:/{name}",
        created_at=completed_at, started_at=completed_at, completed_at=completed_at, file_count=files,
    )
    s.add(op)
    await s.flush()
    return op


async def add_pull(s, user, push, when):
    from models import Operation, OperationStatus, OperationType
    s.add(Operation(
        user_id=user.id, type=OperationType.PULL, status=OperationStatus.COMPLETED,
        source_path=push.dest_path, dest_path=push.source_path, rollback_operation_id=push.id,
        created_at=when, completed_at=when,
    ))
    await s.flush()


async def set_config(s, **values):
    from sqlalchemy import update
    from models import Config
    for key, value in values.items():
        await s.execute(update(Config).where(Config.key == key).values(value=value))


def test_migration_seeds_the_same_settings_as_the_code():
    import importlib.util

    path = Path(__file__).resolve().parent.parent / "alembic" / "versions" / "023_manager_role_and_reports.py"
    spec = importlib.util.spec_from_file_location("revision_023", path)
    revision = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(revision)

    seeded = {key: value for key, value, _type, _description in revision._CONFIG_SEED}
    assert seeded.pop("push_validation_reject_double_dots") == "true"
    assert seeded == report_service.DEFAULTS


@needs_db
async def test_migration_adds_the_manager_role(db):
    from sqlalchemy import select
    from models import User, UserRole

    async with db.session() as s:
        await add_user(s, "grzegorz", "MANAGER")
    async with db.session() as s:
        assert (await s.execute(select(User.role))).scalar() == UserRole.MANAGER


@needs_db
async def test_report_lists_completed_unpulled_pushes_by_warsaw_day(db):
    utc = timezone.utc
    async with db.session() as s:
        natalia, lena = await add_user(s, "mizn"), await add_user(s, "vzle")
        await add_push(s, natalia, "SEPT LAST", datetime(2026, 9, 30, 21, 59, tzinfo=utc))   # 23:59 Warsaw
        await add_push(s, natalia, "OCT FIRST", datetime(2026, 9, 30, 22, 0, tzinfo=utc))    # 00:00 1 Oct
        await add_push(s, natalia, "FAILED", datetime(2026, 10, 2, 10, tzinfo=utc), status="FAILED")
        pulled = await add_push(s, natalia, "PULLED", datetime(2026, 10, 3, 10, tzinfo=utc))
        await add_pull(s, natalia, pulled, datetime(2026, 10, 3, 11, tzinfo=utc))
        await add_push(s, lena, "LENA", datetime(2026, 10, 4, 10, tzinfo=utc))
        await add_push(s, natalia, "OCT LAST", datetime(2026, 10, 31, 22, 59, tzinfo=utc))   # 23:59 31 Oct (CET)
        await add_push(s, natalia, "NOV FIRST", datetime(2026, 10, 31, 23, 0, tzinfo=utc))

    async with db.session() as s:
        entries = await report_service.collect_entries(
            s, ["mizn"], date(2026, 10, 1), date(2026, 11, 1), WARSAW, exclude_pulled=True
        )
        assert [e.catalog_name for e in entries] == ["OCT FIRST", "OCT LAST"]
        with_pulled = await report_service.collect_entries(
            s, ["mizn", "vzle"], date(2026, 10, 1), date(2026, 11, 1), WARSAW, exclude_pulled=False
        )
        assert [e.catalog_name for e in with_pulled] == ["OCT FIRST", "PULLED", "LENA", "OCT LAST"]


@pytest.fixture
def clock(monkeypatch):
    """Run the nightly job as if it were ``when``."""
    async def run_at(when):
        monkeypatch.setattr(report_service, "_now", lambda: when)
        return await report_service.run_scheduled_reports()
    return run_at


@pytest.fixture
def share(monkeypatch):
    """Capture what would be written to the share; ``fail`` makes writes raise."""
    written = []
    state = {"fail": None}

    def write(folder, file_name, data, cfg):
        if state["fail"]:
            raise OSError(state["fail"])
        written.append((folder, file_name, load_workbook(io.BytesIO(data)).active))
        return f"{folder}\\{file_name}"

    monkeypatch.setattr(report_service, "write_to_share", write)
    return written, state


@needs_db
async def test_nightly_job_writes_each_report_once(db, share, clock):
    from sqlalchemy import select
    from models import ReportRun

    run_at = clock
    written, state = share
    utc = timezone.utc
    async with db.session() as s:
        natalia = await add_user(s, "mizn")
        await add_push(s, natalia, "TORBA A", datetime(2026, 10, 31, 12, tzinfo=utc))
        await set_config(s, reports_auto_enabled="true", reports_smb_username="svc",
                         reports_recipients='[{"name": "Natalia", "username": "mizn"}]')

    # Before the run time nothing happens
    assert await run_at(datetime(2026, 11, 1, 0, 30, tzinfo=utc)) == 0
    # 1 Nov 02:30 Warsaw: October is written, once
    first_night = datetime(2026, 11, 1, 1, 30, tzinfo=utc)
    assert await run_at(first_night) == 1
    assert await run_at(first_night + timedelta(minutes=5)) == 0
    folder, name, sheet = written[0]
    assert folder.endswith(r"\2026\10 PAŹDZIERNIK") and name == "Raport_klasyfikacji_Natalia_2026_10.xlsx"
    assert sheet["A3"].value == "TORBA A"

    # 2 Nov: November is started; October finished after it ended, so it is left alone
    state["fail"] = "STATUS_SHARING_VIOLATION (open in Excel)"
    second_night = datetime(2026, 11, 2, 1, 30, tzinfo=utc)
    assert await run_at(second_night) == 1
    async with db.session() as s:
        runs = (await s.execute(
            select(ReportRun).order_by(ReportRun.id)
        )).scalars().all()
        assert [(r.report_month, r.status) for r in runs] == [
            (date(2026, 10, 1), "SUCCESS"), (date(2026, 11, 1), "FAILED"),
        ]
        assert "open in Excel" in runs[1].message

    # The failed one is retried after reports_retry_minutes (60), not before
    state["fail"] = None
    assert await run_at(second_night + timedelta(minutes=30)) == 0
    assert await run_at(second_night + timedelta(minutes=61)) == 1
    assert written[-1][1] == "Raport_klasyfikacji_Natalia_2026_11.xlsx"


@needs_db
async def test_a_month_left_incomplete_is_finished_later(db, share, clock):
    run_at = clock
    written, _ = share
    utc = timezone.utc
    async with db.session() as s:
        await add_user(s, "mizn")
        await set_config(s, reports_auto_enabled="true", reports_smb_username="svc",
                         reports_recipients='[{"name": "Natalia", "username": "mizn"}]')

    # Written on 31 Oct (covering up to 30 Oct), then the server was down on 1 Nov
    assert await run_at(datetime(2026, 10, 31, 2, tzinfo=utc)) == 1
    assert await run_at(datetime(2026, 11, 2, 2, tzinfo=utc)) == 2
    assert [name for _, name, _ in written] == [
        "Raport_klasyfikacji_Natalia_2026_10.xlsx",
        "Raport_klasyfikacji_Natalia_2026_11.xlsx",
        "Raport_klasyfikacji_Natalia_2026_10.xlsx",
    ]
    # A month RFM never wrote (September) is not touched
    assert not any(name.endswith("2026_09.xlsx") for _, name, _ in written)


@needs_db
async def test_reports_stay_off_by_default(db, share, clock):
    run_at = clock
    written, _ = share
    assert await run_at(datetime(2026, 11, 1, 3, tzinfo=timezone.utc)) == 0
    assert written == []


@needs_db
async def test_manager_history_filters_and_summary(db):
    from api.routes.manager import search_operations

    utc = timezone.utc
    async with db.session() as s:
        manager = await add_user(s, "grzegorz", "MANAGER")
        natalia, lena = await add_user(s, "mizn"), await add_user(s, "vzle")
        await add_push(s, natalia, "TORBA HB0788", datetime(2026, 10, 1, 10, tzinfo=utc), files=4)
        pulled = await add_push(s, natalia, "BUTY FU1479", datetime(2026, 10, 2, 10, tzinfo=utc))
        await add_pull(s, natalia, pulled, datetime(2026, 10, 2, 11, tzinfo=utc))
        await add_push(s, lena, "BUTY CARTER", datetime(2026, 10, 3, 10, tzinfo=utc), status="FAILED")
        await add_push(s, lena, "SZALIK", datetime(2026, 9, 3, 10, tzinfo=utc))

    async with db.session() as s:
        october = dict(date_from=datetime(2026, 10, 1, tzinfo=WARSAW), date_to=datetime(2026, 11, 1, tzinfo=WARSAW))
        everything = await search_operations(manager, s, **october, limit=50, offset=0, q=None, users=None,
                                             types=None, statuses=None, pulled=None, sort_by="created_at",
                                             sort_order="asc")
        assert [op["catalog_name"] for op in everything["operations"]] == [
            "TORBA HB0788", "BUTY FU1479", "BUTY FU1479", "BUTY CARTER",
        ]
        summary = everything["summary"]
        assert summary["total"] == 4
        assert summary["by_type"] == {"PUSH": 3, "PULL": 1}
        assert summary["by_status"] == {"COMPLETED": 3, "FAILED": 1}
        assert summary["by_user"][0] == {
            "username": "mizn", "total": 3, "push": 2, "pull": 1, "update": 0, "failed": 0, "files": 7,
        }

        buty_still_published = await search_operations(
            manager, s, **october, q="buty", users="mizn,vzle", types="push", statuses=None, pulled=False,
            sort_by="created_at", sort_order="desc", limit=50, offset=0,
        )
        assert [op["catalog_name"] for op in buty_still_published["operations"]] == ["BUTY CARTER"]

        pulled_only = await search_operations(
            manager, s, **october, q=None, users=None, types=None, statuses=None, pulled=True,
            sort_by="created_at", sort_order="desc", limit=50, offset=0,
        )
        assert [(op["catalog_name"], op["has_been_pulled"]) for op in pulled_only["operations"]] == [
            ("BUTY FU1479", True),
        ]


async def test_role_levels():
    from api.middleware.auth import PermissionDeniedError, require_admin, require_manager, require_user
    from models import User, UserRole

    def as_role(role):
        return User(username=role.value, password_hash="x", role=role)

    for role in UserRole:
        assert await require_user(as_role(role)) is not None
    assert await require_manager(as_role(UserRole.MANAGER)) is not None
    assert await require_manager(as_role(UserRole.ADMIN)) is not None
    with pytest.raises(PermissionDeniedError):
        await require_manager(as_role(UserRole.USER))
    with pytest.raises(PermissionDeniedError):
        await require_admin(as_role(UserRole.MANAGER))
