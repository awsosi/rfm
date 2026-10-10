"""
Classification reports: the catalogs a person pushed, day by day, as XLSX.

Asked for by Grzegorz Gutek on 2026-10-07: one workbook per person and month,
``Raport_klasyfikacji_<name>_<year>_<month>.xlsx``, in a folder per month on a
Windows share. Every setting is a ``reports_*`` row of the config table
(Admin Panel -> Configuration -> Classification Reports).

Layout of a report:

- a header row with the configured column headers
- per day with pushes, a merged row holding the date (``reports_day_row_color``)
  followed by one row per catalog pushed that day
- a catalog pushed after the date in ``reports_late_field`` (by default the
  product's release date) gets ``reports_late_row_color``

Product columns (``product.<key>``) come from the PolkaSQL web service
``RFM_ProductDetails`` (docs/polkasql/RFM_ProductDetails.sql). Without it they
stay blank and the report is still written, with a warning on the run.

The nightly job (``run_scheduled_reports``, every API process calls it once a
minute) rewrites the whole report of the month that "yesterday" belongs to, so
a missed night heals itself and the run on the 1st completes the month before.
Each scheduled report is one ``report_runs`` row, unique per local run date,
recipient and month: inserting it is how a single process claims the work.
A failed or stuck run is claimed again after ``reports_retry_minutes``.

Files go to the share over SMB (``smbprotocol``, pure Python, from the API
container). A file is written next to the target and then renamed over it, so
a reader never sees half a workbook; a report open in Excel cannot be replaced
and the run fails until the next retry. Unless ``reports_overwrite_foreign``
is on, an existing file that RFM did not create is never overwritten.
"""

import asyncio
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from itertools import groupby
from typing import Optional
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import httpx
from loguru import logger
from sqlalchemy import and_, exists, or_, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import aliased

from api.services.polka import POLKA_HEADERS, polka_json
from database import DatabaseManager
from models import (
    Config,
    Operation,
    OperationStatus,
    OperationType,
    ReportRun,
    ReportRunStatus,
    User,
)

# Written into the workbook's properties; a file without it was not made by RFM
REPORT_CREATOR = "RFM"

# Fields a column can show besides product.<key>
BUILTIN_FIELDS = ("catalog_name", "pushed_at", "pushed_date", "username", "file_count", "operation_id")

TEMPLATE_PLACEHOLDERS = ("{year}", "{month}", "{month_name}", "{name}", "{username}")

# Concurrent requests to the product details service
_PRODUCT_CONCURRENCY = 4

# Same defaults as alembic revision 023; used when a row is missing
DEFAULTS = {
    "reports_auto_enabled": "false",
    "reports_run_time": "02:00",
    "reports_timezone": "Europe/Warsaw",
    "reports_retry_minutes": "60",
    "reports_recipients": json.dumps([
        {"name": "Natalia", "username": "mizn"},
        {"name": "Lena", "username": "vzle"},
        {"name": "Ewa", "username": "vzwe"},
    ], ensure_ascii=False),
    "reports_output_dir": (
        "\\\\radius1\\Users\\Iza.Horna\\WAŻNE FOLDERY BEATA GRZESIEK\\"
        "!RAPORTY KLASYFIKACJA\\{year}\\{month} {month_name}"
    ),
    "reports_file_name": "Raport_klasyfikacji_{name}_{year}_{month}.xlsx",
    "reports_month_names": (
        "STYCZEŃ,LUTY,MARZEC,KWIECIEŃ,MAJ,CZERWIEC,LIPIEC,SIERPIEŃ,WRZESIEŃ,"
        "PAŹDZIERNIK,LISTOPAD,GRUDZIEŃ"
    ),
    "reports_smb_username": "",
    "reports_smb_password": "",
    "reports_columns": json.dumps([
        {"header": "Nazwa", "field": "catalog_name"},
        {"header": "Projektant", "field": "product.designer"},
        {"header": "Płeć", "field": "product.gender"},
        {"header": "Data premiery", "field": "product.release_date"},
        {"header": "Data ostatniej wysyłki", "field": "product.last_send_date"},
    ], ensure_ascii=False),
    "reports_sheet_name": "Raport",
    "reports_date_format": "DD.MM.YYYY",
    "reports_day_row_color": "#afd095",
    "reports_late_row_color": "#ffa6a6",
    "reports_late_field": "product.release_date",
    "reports_exclude_pulled": "true",
    "reports_overwrite_foreign": "false",
    "reports_product_url": "",
    "reports_product_api_key": "",
    "reports_product_timeout": "20",
}

_COLOR_RE = re.compile(r"#?([0-9A-Fa-f]{6})")
_TIME_RE = re.compile(r"([01]?\d|2[0-3]):([0-5]\d)")
_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")
_DATETIME_RE = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?")
# Characters Windows refuses in a file or folder name
_UNSAFE_NAME_RE = re.compile(r'[\\/:*?"<>|]')


def _now() -> datetime:
    """The clock of runs and the schedule (tests replace it)."""
    return datetime.now(timezone.utc)


class ReportConfigError(ValueError):
    """A reports_* setting cannot be used; the message names the setting."""


@dataclass
class ReportConfig:
    auto_enabled: bool
    run_time: time
    tz: ZoneInfo
    retry_minutes: int
    recipients: list
    output_dir: str
    file_name: str
    month_names: list
    smb_username: str
    smb_password: str
    columns: list
    sheet_name: str
    date_format: str
    day_row_color: str
    late_row_color: str
    late_field: str
    exclude_pulled: bool
    overwrite_foreign: bool
    product_url: str
    product_api_key: str
    product_timeout: int

    def public(self) -> dict:
        """What a manager may see: everything but the share password and API key."""
        return {
            "auto_enabled": self.auto_enabled,
            "run_time": self.run_time.strftime("%H:%M"),
            "timezone": str(self.tz),
            "retry_minutes": self.retry_minutes,
            "recipients": self.recipients,
            "output_dir": self.output_dir,
            "file_name": self.file_name,
            "columns": self.columns,
            "late_field": self.late_field,
            "exclude_pulled": self.exclude_pulled,
            "product_data": bool(self.product_url and self.product_api_key),
            "share_account": bool(self.smb_username),
        }


@dataclass
class ReportEntry:
    operation_id: int
    catalog_name: str
    pushed_at: datetime  # aware, UTC
    username: str
    file_count: Optional[int]
    product: dict = field(default_factory=dict)


def _truthy(value: Optional[str]) -> bool:
    return (value or "").strip().lower() in ("true", "1", "yes")


def _color(value: str, key: str) -> str:
    match = _COLOR_RE.fullmatch((value or "").strip())
    if not match:
        raise ReportConfigError(f"{key} must be a colour like #afd095, not {value!r}")
    return match.group(1).upper()


def _json_list(value: str, key: str) -> list:
    try:
        parsed = json.loads(value or "[]")
    except json.JSONDecodeError as exc:
        raise ReportConfigError(f"{key} is not valid JSON: {exc}") from exc
    if not isinstance(parsed, list):
        raise ReportConfigError(f"{key} must be a JSON list")
    return parsed


def parse_report_config(values: dict) -> ReportConfig:
    """Build a ``ReportConfig`` from config rows; missing rows take ``DEFAULTS``."""
    get = lambda key: values.get(key, DEFAULTS[key])  # noqa: E731

    run_time = _TIME_RE.fullmatch(get("reports_run_time").strip())
    if not run_time:
        raise ReportConfigError(f"reports_run_time must be HH:MM, not {get('reports_run_time')!r}")
    try:
        tz = ZoneInfo(get("reports_timezone").strip())
    except (ZoneInfoNotFoundError, ValueError) as exc:
        raise ReportConfigError(f"reports_timezone {get('reports_timezone')!r} is not a known time zone") from exc

    try:
        retry_minutes = max(int(get("reports_retry_minutes")), 1)
        product_timeout = max(int(get("reports_product_timeout")), 1)
    except ValueError as exc:
        raise ReportConfigError("reports_retry_minutes and reports_product_timeout must be numbers") from exc

    recipients = []
    for item in _json_list(get("reports_recipients"), "reports_recipients"):
        if not isinstance(item, dict) or not str(item.get("username") or "").strip():
            raise ReportConfigError('each reports_recipients entry needs a "username"')
        username = str(item["username"]).strip()
        recipients.append({"name": str(item.get("name") or "").strip() or username, "username": username})

    columns = []
    for item in _json_list(get("reports_columns"), "reports_columns"):
        if not isinstance(item, dict) or not str(item.get("field") or "").strip():
            raise ReportConfigError('each reports_columns entry needs a "field"')
        name = str(item["field"]).strip()
        if name not in BUILTIN_FIELDS and not (name.startswith("product.") and len(name) > len("product.")):
            raise ReportConfigError(
                f"reports_columns field {name!r} is unknown; use {', '.join(BUILTIN_FIELDS)} or product.<key>"
            )
        columns.append({"header": str(item.get("header") or name), "field": name})
    if not columns:
        raise ReportConfigError("reports_columns must list at least one column")

    month_names = [m.strip() for m in get("reports_month_names").split(",")]
    if len(month_names) != 12 or not all(month_names):
        raise ReportConfigError("reports_month_names must hold twelve comma-separated names")

    sheet_name = _UNSAFE_NAME_RE.sub("_", get("reports_sheet_name").strip()).replace("[", "(").replace("]", ")")

    return ReportConfig(
        auto_enabled=_truthy(get("reports_auto_enabled")),
        run_time=time(int(run_time.group(1)), int(run_time.group(2))),
        tz=tz,
        retry_minutes=retry_minutes,
        recipients=recipients,
        output_dir=get("reports_output_dir").strip(),
        file_name=get("reports_file_name").strip(),
        month_names=month_names,
        smb_username=get("reports_smb_username").strip(),
        smb_password=get("reports_smb_password"),
        columns=columns,
        sheet_name=(sheet_name or "Raport")[:31],
        date_format=get("reports_date_format").strip() or DEFAULTS["reports_date_format"],
        day_row_color=_color(get("reports_day_row_color"), "reports_day_row_color"),
        late_row_color=_color(get("reports_late_row_color"), "reports_late_row_color"),
        late_field=get("reports_late_field").strip(),
        exclude_pulled=_truthy(get("reports_exclude_pulled")),
        overwrite_foreign=_truthy(get("reports_overwrite_foreign")),
        product_url=get("reports_product_url").strip(),
        product_api_key=get("reports_product_api_key").strip(),
        product_timeout=product_timeout,
    )


async def load_report_config(db) -> ReportConfig:
    result = await db.execute(select(Config.key, Config.value).where(Config.key.in_(DEFAULTS)))
    return parse_report_config({row.key: row.value for row in result})


# ---------------------------------------------------------------------------
# Periods and paths
# ---------------------------------------------------------------------------

def month_start(day: date) -> date:
    return day.replace(day=1)


def next_month(first: date) -> date:
    return (first.replace(day=28) + timedelta(days=4)).replace(day=1)


def _local_midnight_utc(day: date, tz: ZoneInfo) -> datetime:
    return datetime.combine(day, time(0), tzinfo=tz).astimezone(timezone.utc)


def render_template(template: str, config: ReportConfig, month: date, name: str, username: str) -> str:
    """Fill the path/file name placeholders; values cannot add folders or forbidden characters."""
    values = {
        "{year}": f"{month.year:04d}",
        "{month}": f"{month.month:02d}",
        "{month_name}": config.month_names[month.month - 1],
        "{name}": name,
        "{username}": username,
    }
    for placeholder, value in values.items():
        template = template.replace(placeholder, _UNSAFE_NAME_RE.sub("_", value))
    return template


def report_location(config: ReportConfig, month: date, recipient: dict) -> tuple:
    """``(folder, file name)`` of a recipient's monthly report on the share."""
    folder = render_template(config.output_dir, config, month, recipient["name"], recipient["username"])
    folder = folder.replace("/", "\\").rstrip("\\")
    if not folder.startswith("\\\\") or len(folder.split("\\")) < 4:
        raise ReportConfigError(f"reports_output_dir must be a UNC path like \\\\server\\share\\folder, not {folder!r}")
    file_name = render_template(config.file_name, config, month, recipient["name"], recipient["username"])
    file_name = _UNSAFE_NAME_RE.sub("_", file_name)
    if not file_name.lower().endswith(".xlsx"):
        file_name += ".xlsx"
    return folder, file_name


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------

async def collect_entries(
    db, usernames: list, start: date, end: date, tz: ZoneInfo, exclude_pulled: bool
) -> list:
    """
    Completed PUSHes of ``usernames`` whose local completion day is in
    [``start``, ``end``), oldest first.
    """
    from api.services.pim_service import extract_catalog_name

    pushed_at = Operation.completed_at
    stmt = (
        select(Operation, User.username)
        .join(User, Operation.user_id == User.id)
        .where(
            Operation.type == OperationType.PUSH,
            Operation.status == OperationStatus.COMPLETED,
            User.username.in_(usernames),
            pushed_at >= _local_midnight_utc(start, tz),
            pushed_at < _local_midnight_utc(end, tz),
        )
        .order_by(pushed_at, Operation.id)
    )
    if exclude_pulled:
        pull = aliased(Operation)
        stmt = stmt.where(~exists().where(
            pull.rollback_operation_id == Operation.id,
            pull.type == OperationType.PULL,
            pull.status == OperationStatus.COMPLETED,
        ))

    entries = []
    for operation, username in (await db.execute(stmt)).all():
        entries.append(ReportEntry(
            operation_id=operation.id,
            catalog_name=extract_catalog_name("PUSH", operation.source_path, operation.dest_path),
            pushed_at=operation.completed_at,
            username=username,
            file_count=operation.file_count,
        ))
    return entries


async def fetch_products(names: list, config: ReportConfig) -> tuple:
    """
    Product details per catalog name from ``RFM_ProductDetails``.

    Returns ``(products, warning)``: products maps a name to its fields (empty
    when unknown); warning is None when every name was found.
    """
    unique = sorted(set(names))
    if not unique:
        return {}, None
    if not config.product_url or not config.product_api_key:
        return {}, "Product details service is not configured (reports_product_url, reports_product_api_key)"

    semaphore = asyncio.Semaphore(_PRODUCT_CONCURRENCY)
    products, problems = {}, []

    async def lookup(client, name):
        async with semaphore:
            try:
                response = await client.get(
                    config.product_url,
                    params={"ApiKey": config.product_api_key, "CatalogName": name},
                    headers=POLKA_HEADERS,
                )
                if response.status_code != 200:
                    problems.append(f"{name}: HTTP {response.status_code}")
                    return
                data = polka_json(response)
            except Exception as exc:
                problems.append(f"{name}: {type(exc).__name__}: {exc}")
                return
        if not data.get("success"):
            problems.append(f"{name}: {data.get('error') or 'service error'}")
        elif not data.get("found") or not isinstance(data.get("product"), dict):
            problems.append(f"{name}: not found")
        else:
            products[name] = data["product"]

    async with httpx.AsyncClient(timeout=config.product_timeout) as client:
        await asyncio.gather(*(lookup(client, name) for name in unique))

    warning = None
    if problems:
        problems.sort()
        warning = (
            f"No product data for {len(problems)} of {len(unique)} catalog(s): "
            + "; ".join(problems[:5]) + ("; ..." if len(problems) > 5 else "")
        )
    return products, warning


def _coerce(value):
    """Dates arrive from PolkaSQL as text; make them real Excel dates."""
    if isinstance(value, str):
        text = value.strip()
        if _DATE_RE.fullmatch(text):
            return date.fromisoformat(text)
        if _DATETIME_RE.match(text):
            stamp = datetime.fromisoformat(text[:19].replace("T", " "))
            return stamp.date() if stamp.time() == time(0) else stamp
        return text
    return value


def field_value(entry: ReportEntry, name: str, tz: ZoneInfo):
    if name.startswith("product."):
        return _coerce(entry.product.get(name[len("product."):]))
    local = entry.pushed_at.astimezone(tz)
    return {
        "catalog_name": entry.catalog_name,
        "pushed_at": local.replace(tzinfo=None, microsecond=0),
        "pushed_date": local.date(),
        "username": entry.username,
        "file_count": entry.file_count,
        "operation_id": entry.operation_id,
    }.get(name)


def is_late(entry: ReportEntry, config: ReportConfig) -> bool:
    """Pushed on a later day than the date in ``reports_late_field``."""
    if not config.late_field:
        return False
    limit = field_value(entry, config.late_field, config.tz)
    if isinstance(limit, datetime):
        limit = limit.date()
    if not isinstance(limit, date):
        return False
    return entry.pushed_at.astimezone(config.tz).date() > limit


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def build_workbook(entries: list, config: ReportConfig, columns: Optional[list] = None) -> bytes:
    """The report as XLSX bytes; ``columns`` overrides ``config.columns``."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    columns = columns or config.columns
    width = len(columns)
    workbook = Workbook()
    workbook.properties.creator = REPORT_CREATOR
    sheet = workbook.active
    sheet.title = config.sheet_name

    bold = Font(bold=True)
    day_fill = PatternFill("solid", fgColor=config.day_row_color)
    late_fill = PatternFill("solid", fgColor=config.late_row_color)
    thin = Side(style="thin", color="999999")
    datetime_format = config.date_format + " hh:mm"
    widths = [len(str(column["header"])) for column in columns]

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column["header"])
        cell.font = bold
        cell.border = Border(bottom=thin)
    sheet.freeze_panes = "A2"

    row = 2
    day_of = lambda entry: entry.pushed_at.astimezone(config.tz).date()  # noqa: E731
    for day, day_entries in groupby(entries, key=day_of):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        cell = sheet.cell(row=row, column=1, value=day)
        cell.number_format = config.date_format
        cell.font = bold
        cell.alignment = Alignment(horizontal="left")
        for index in range(1, width + 1):
            sheet.cell(row=row, column=index).fill = day_fill
        row += 1

        for entry in day_entries:
            late = is_late(entry, config)
            for index, column in enumerate(columns, start=1):
                value = field_value(entry, column["field"], config.tz)
                cell = sheet.cell(row=row, column=index, value=value)
                if isinstance(value, datetime):
                    cell.number_format = datetime_format
                elif isinstance(value, date):
                    cell.number_format = config.date_format
                if late:
                    cell.fill = late_fill
                shown = len(config.date_format) if isinstance(value, date) else len(str(value or ""))
                widths[index - 1] = max(widths[index - 1], shown)
            row += 1

    for index, used in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(used + 2, 10), 70)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def generate_report(
    db, config: ReportConfig, usernames: list, start: date, end: date, columns: Optional[list] = None
) -> tuple:
    """``(xlsx bytes, number of catalogs, warning)`` for pushes in [start, end)."""
    entries = await collect_entries(db, usernames, start, end, config.tz, config.exclude_pulled)
    warning = None
    if any(c["field"].startswith("product.") for c in (columns or config.columns)) or config.late_field.startswith("product."):
        products, warning = await fetch_products([e.catalog_name for e in entries], config)
        for entry in entries:
            entry.product = products.get(entry.catalog_name, {})
    return build_workbook(entries, config, columns), len(entries), warning


# ---------------------------------------------------------------------------
# Windows share (SMB)
# ---------------------------------------------------------------------------

def created_by_rfm(data: bytes) -> bool:
    """Whether an XLSX carries RFM as its creator (kept when Excel saves it again)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            core = archive.read("docProps/core.xml").decode("utf-8", "replace")
    except (zipfile.BadZipFile, KeyError):
        return False
    match = re.search(r"<dc:creator>([^<]*)</dc:creator>", core)
    return bool(match and match.group(1).strip() == REPORT_CREATOR)


def _smb_credentials(config: ReportConfig) -> dict:
    if not config.smb_username:
        raise ReportConfigError("reports_smb_username is empty: set the account that writes to the share")
    return {"username": config.smb_username, "password": config.smb_password}


def write_to_share(folder: str, file_name: str, data: bytes, config: ReportConfig) -> str:
    """Write ``data`` to ``folder\\file_name`` (blocking: run it in a thread). Returns the path."""
    import smbclient

    credentials = _smb_credentials(config)
    target = f"{folder}\\{file_name}"
    smbclient.makedirs(folder, exist_ok=True, **credentials)

    if not config.overwrite_foreign:
        try:
            with smbclient.open_file(target, mode="rb", **credentials) as existing:
                current = existing.read()
        except FileNotFoundError:
            current = None
        if current is not None and not created_by_rfm(current):
            raise RuntimeError(
                f"{target} exists and was not created by RFM; move it away or turn on reports_overwrite_foreign"
            )

    partial = f"{folder}\\~{file_name}.rfm-partial"
    with smbclient.open_file(partial, mode="wb", **credentials) as handle:
        handle.write(data)
    try:
        smbclient.replace(partial, target, **credentials)
    except Exception:
        try:
            smbclient.remove(partial, **credentials)
        except Exception:
            pass
        raise
    return target


def check_share(config: ReportConfig) -> str:
    """Create the current month's folder and a test file, then remove the file. Returns the folder."""
    import smbclient

    credentials = _smb_credentials(config)
    today = datetime.now(config.tz).date()
    recipient = (config.recipients or [{"name": "test", "username": "test"}])[0]
    folder, _ = report_location(config, month_start(today), recipient)
    smbclient.makedirs(folder, exist_ok=True, **credentials)
    probe = f"{folder}\\~rfm-write-test.tmp"
    with smbclient.open_file(probe, mode="wb", **credentials) as handle:
        handle.write(b"RFM write test")
    smbclient.remove(probe, **credentials)
    return folder


# ---------------------------------------------------------------------------
# Runs
# ---------------------------------------------------------------------------

async def _finish(run_id: int, **values) -> None:
    async with DatabaseManager.session() as db:
        await db.execute(
            update(ReportRun).where(ReportRun.id == run_id).values(finished_at=_now(), **values)
        )


async def execute_run(run_id: int) -> None:
    """Generate one recipient's monthly report and write it to the share; records the outcome."""
    try:
        async with DatabaseManager.session() as db:
            run = await db.get(ReportRun, run_id)
            config = await load_report_config(db)
            recipient = {"name": run.label, "username": run.recipient}
            month = run.report_month
            folder, file_name = report_location(config, month, recipient)
            data, rows, warning = await generate_report(
                db, config, [run.recipient], month, next_month(month)
            )
        path = await asyncio.to_thread(write_to_share, folder, file_name, data, config)
    except Exception as exc:
        logger.error(f"Report run {run_id} failed: {type(exc).__name__}: {exc}")
        await _finish(run_id, status=ReportRunStatus.FAILED, message=f"{type(exc).__name__}: {exc}"[:2000])
        return
    logger.info(f"Report run {run_id}: {rows} catalog(s) written to {path}")
    await _finish(run_id, status=ReportRunStatus.SUCCESS, file_path=path, row_count=rows, message=warning)


async def start_manual_runs(recipients: list, month: date, requested_by: str) -> list:
    """Insert MANUAL runs for ``recipients`` and start them in the background; returns the run ids."""
    async with DatabaseManager.session() as db:
        runs = [
            ReportRun(
                trigger="MANUAL", report_month=month, recipient=r["username"], label=r["name"],
                status=ReportRunStatus.RUNNING, requested_by=requested_by, started_at=_now(),
            )
            for r in recipients
        ]
        db.add_all(runs)
        await db.flush()
        ids = [run.id for run in runs]
    for run_id in ids:
        _spawn(execute_run(run_id))
    return ids


_background: set = set()


def _spawn(coroutine) -> None:
    task = asyncio.create_task(coroutine)
    _background.add(task)
    task.add_done_callback(_background.discard)


async def _claim_scheduled(run_date: date, recipient: dict, month: date, retry_minutes: int) -> Optional[int]:
    """
    Claim the scheduled report of ``recipient`` for ``month`` on ``run_date``:
    insert its row, or take over a FAILED one (or one stuck RUNNING) once
    ``retry_minutes`` have passed. Returns the run id, or None when another
    process has it or it is done.
    """
    now = _now()
    async with DatabaseManager.session() as db:
        inserted = await db.execute(
            pg_insert(ReportRun)
            .values(
                trigger="SCHEDULED", run_date=run_date, report_month=month,
                recipient=recipient["username"], label=recipient["name"],
                status=ReportRunStatus.RUNNING, attempts=1, started_at=now,
            )
            .on_conflict_do_nothing(
                index_elements=["run_date", "recipient", "report_month"],
                index_where=ReportRun.trigger == "SCHEDULED",
            )
            .returning(ReportRun.id)
        )
        run_id = inserted.scalar()
        if run_id is not None:
            return run_id

        cutoff = now - timedelta(minutes=retry_minutes)
        retried = await db.execute(
            update(ReportRun)
            .where(
                ReportRun.trigger == "SCHEDULED",
                ReportRun.run_date == run_date,
                ReportRun.recipient == recipient["username"],
                ReportRun.report_month == month,
                or_(
                    and_(ReportRun.status == ReportRunStatus.FAILED, ReportRun.finished_at <= cutoff),
                    and_(ReportRun.status == ReportRunStatus.RUNNING, ReportRun.started_at <= cutoff),
                ),
            )
            .values(
                status=ReportRunStatus.RUNNING, attempts=ReportRun.attempts + 1, label=recipient["name"],
                started_at=now, finished_at=None, message=None,
            )
            .returning(ReportRun.id)
        )
        return retried.scalar()


async def _needs_completion(db, recipient: dict, month: date, tz: ZoneInfo) -> bool:
    """
    Whether RFM reported ``month`` for ``recipient`` before but never after the
    month ended, i.e. its last days are missing. A month RFM never touched is
    left alone, so turning reports on does not overwrite older files.
    """
    base = select(ReportRun.id).where(
        ReportRun.recipient == recipient["username"], ReportRun.report_month == month
    )
    touched = (await db.execute(base.limit(1))).first() is not None
    if not touched:
        return False
    completed = (await db.execute(base.where(
        ReportRun.status == ReportRunStatus.SUCCESS,
        ReportRun.finished_at >= _local_midnight_utc(next_month(month), tz),
    ).limit(1))).first() is not None
    return not completed


async def run_scheduled_reports() -> int:
    """
    Nightly job, called every minute by each API process. After
    ``reports_run_time`` it writes, for every recipient, the report of the
    month that yesterday belongs to (and completes the month before when its
    last days are missing). Returns the number of runs this process executed.
    """
    async with DatabaseManager.session() as db:
        try:
            config = await load_report_config(db)
        except ReportConfigError as exc:
            if _truthy((await db.execute(
                select(Config.value).where(Config.key == "reports_auto_enabled")
            )).scalar()):
                logger.error(f"Classification reports are on but cannot run: {exc}")
            return 0
        if not config.auto_enabled or not config.recipients:
            return 0

        local_now = _now().astimezone(config.tz)
        if local_now.time() < config.run_time:
            return 0
        today = local_now.date()
        current = month_start(today - timedelta(days=1))
        previous = month_start(current - timedelta(days=1))

        work = []
        for recipient in config.recipients:
            work.append((recipient, current))
            if await _needs_completion(db, recipient, previous, config.tz):
                work.append((recipient, previous))

    executed = 0
    for recipient, month in work:
        run_id = await _claim_scheduled(today, recipient, month, config.retry_minutes)
        if run_id is not None:
            await execute_run(run_id)
            executed += 1
    return executed
