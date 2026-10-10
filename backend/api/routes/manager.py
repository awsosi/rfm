"""
Manager view: operation history with precise filters, and classification reports.

Every endpoint needs the MANAGER role or higher (admins included). Managers
only read: they never change users, workers or settings. The one thing they
start is a report run, which writes the configured files to the share.
"""

import io
import re
from datetime import date, datetime, timedelta, timezone
from typing import Annotated, List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import case, desc, exists, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import aliased

from api.middleware.auth import require_manager
from api.middleware.logging import AuditLogger, get_client_ip
from api.schemas import OperationResponse
from api.services.elasticsearch_service import search_terms
from api.services.operation_service import get_integration_status, get_update_ids_for_pushes
from api.services.pim_service import extract_catalog_name
from api.services import report_service
from database import get_db
from models import Operation, OperationStatus, OperationType, ReportRun, User

router = APIRouter(prefix="/api/manager", tags=["manager"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Largest history export: one click cannot build an unbounded workbook, and the
# id lists of the follow-up queries stay under asyncpg's 32767 bind parameters
EXPORT_LIMIT = 20000

_SORT_COLUMNS = {
    "id": Operation.id,
    "created_at": Operation.created_at,
    "completed_at": Operation.completed_at,
    "type": Operation.type,
    "status": Operation.status,
    "user": User.username,
    "catalog": Operation.source_path,
    "file_count": Operation.file_count,
}


# =============================================================================
# Schemas
# =============================================================================


class ManagerUser(BaseModel):
    id: int
    username: str
    role: str
    is_active: bool


class OperationFilters(BaseModel):
    """History filters; every given filter must match."""

    q: Optional[str] = None
    users: List[str] = Field(default_factory=list)
    types: List[OperationType] = Field(default_factory=list)
    statuses: List[OperationStatus] = Field(default_factory=list)
    # Timestamps (with offset); created_at in [date_from, date_to)
    date_from: Optional[datetime] = None
    date_to: Optional[datetime] = None
    # PUSH only: True = pulled back later, False = still published
    pulled: Optional[bool] = None


class ExportColumn(BaseModel):
    key: str
    header: str


class ExportRequest(OperationFilters):
    sort_by: str = "created_at"
    sort_order: str = "desc"
    # Headers come from the browser so the file speaks the user's language
    columns: List[ExportColumn]


class ReportDownloadRequest(BaseModel):
    usernames: List[str] = Field(min_length=1)
    # Local days (reports_timezone), both included
    date_from: date
    date_to: date
    include_user_column: bool = False
    user_column_header: str = "User"


class ReportRunRequest(BaseModel):
    month: str = Field(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")
    # Configured recipients to write; empty = all of them
    usernames: List[str] = Field(default_factory=list)


class ReportRunResponse(BaseModel):
    id: int
    trigger: str
    run_date: Optional[date]
    report_month: date
    recipient: str
    label: str
    status: str
    attempts: int
    file_path: Optional[str]
    row_count: Optional[int]
    message: Optional[str]
    requested_by: Optional[str]
    started_at: datetime
    finished_at: Optional[datetime]

    model_config = {"from_attributes": True}


# =============================================================================
# Users
# =============================================================================


@router.get("/users", response_model=List[ManagerUser])
async def list_users(
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Everyone who can appear in the history, for the user filter."""
    users = (await db.execute(select(User).order_by(User.username))).scalars()
    return [
        ManagerUser(id=u.id, username=u.username, role=u.role.value, is_active=u.is_active)
        for u in users
    ]


# =============================================================================
# Operation history
# =============================================================================


def _pulled_clause():
    pull = aliased(Operation)
    return exists().where(
        pull.rollback_operation_id == Operation.id,
        pull.type == OperationType.PULL,
        pull.status == OperationStatus.COMPLETED,
    )


def filtered_operations(filters: OperationFilters):
    """``select(Operation, User)`` narrowed by ``filters``, unsorted and unpaginated."""
    query = select(Operation, User).join(User, Operation.user_id == User.id)
    if filters.users:
        query = query.where(User.username.in_(filters.users))
    if filters.types:
        query = query.where(Operation.type.in_(filters.types))
    if filters.statuses:
        query = query.where(Operation.status.in_(filters.statuses))
    if filters.date_from:
        query = query.where(Operation.created_at >= filters.date_from)
    if filters.date_to:
        query = query.where(Operation.created_at < filters.date_to)
    if filters.pulled is not None:
        query = query.where(Operation.type == OperationType.PUSH)
        query = query.where(_pulled_clause() if filters.pulled else ~_pulled_clause())
    # Same rule as the explorer's history search: every word somewhere
    for term in search_terms(filters.q):
        pattern = "%" + re.sub(r"[\\%_]", r"\\\g<0>", term) + "%"
        query = query.where(or_(*(
            column.ilike(pattern, escape="\\")
            for column in (
                Operation.source_path,
                Operation.dest_path,
                Operation.original_path,
                Operation.archive_path,
                Operation.error_msg,
                User.username,
            )
        )))
    return query


def _sorted(query, sort_by: str, sort_order: str):
    column = _SORT_COLUMNS.get(sort_by, Operation.created_at)
    ordered = desc(column).nulls_last() if sort_order.lower() != "asc" else column.asc().nulls_last()
    return query.order_by(ordered, desc(Operation.id))


async def _responses(db: AsyncSession, rows) -> list:
    """History rows as OperationResponse dicts, with pull/UPDATE/PIM state and the catalog name."""
    ids = [op.id for op, _ in rows]
    push_ids = [op.id for op, _ in rows if op.type == OperationType.PUSH and op.status == OperationStatus.COMPLETED]
    pulled = set()
    if push_ids:
        pull = aliased(Operation)
        pulled = {
            row[0] for row in await db.execute(
                select(pull.rollback_operation_id).where(
                    pull.rollback_operation_id.in_(push_ids),
                    pull.type == OperationType.PULL,
                    pull.status == OperationStatus.COMPLETED,
                )
            )
        }
    update_ids = await get_update_ids_for_pushes(db, push_ids)
    integration = await get_integration_status(db, ids)

    responses = []
    for operation, user in rows:
        response = OperationResponse.model_validate(operation).model_copy(update={
            "user_name": user.username,
            "has_been_pulled": operation.id in pulled,
            "update_operation_ids": update_ids.get(operation.id, []),
            **integration.get(operation.id, {}),
        })
        item = response.model_dump(mode="json")
        item["catalog_name"] = extract_catalog_name(
            operation.type.value, operation.source_path, operation.dest_path
        )
        responses.append(item)
    return responses


async def _summary(db: AsyncSession, filters: OperationFilters) -> dict:
    """Counts by type, status and user for everything the filters match."""
    matched = filtered_operations(filters).with_only_columns(
        Operation.id, Operation.type, Operation.status, Operation.file_count, User.username
    ).subquery()

    by_type = {row[0].value: row[1] for row in await db.execute(
        select(matched.c.type, func.count()).group_by(matched.c.type)
    )}
    by_status = {row[0].value: row[1] for row in await db.execute(
        select(matched.c.status, func.count()).group_by(matched.c.status)
    )}
    count_of = lambda condition: func.count(case((condition, 1)))  # noqa: E731
    by_user = [
        {
            "username": row.username,
            "total": row.total,
            "push": row.push,
            "pull": row.pull,
            "update": row.update,
            "failed": row.failed,
            "files": int(row.files or 0),
        }
        for row in await db.execute(
            select(
                matched.c.username,
                func.count().label("total"),
                count_of(matched.c.type == OperationType.PUSH).label("push"),
                count_of(matched.c.type == OperationType.PULL).label("pull"),
                count_of(matched.c.type == OperationType.UPDATE).label("update"),
                count_of(matched.c.status == OperationStatus.FAILED).label("failed"),
                func.sum(case((matched.c.type == OperationType.PUSH, matched.c.file_count), else_=0)).label("files"),
            )
            .group_by(matched.c.username)
            .order_by(desc("total"), matched.c.username)
        )
    ]
    return {
        "total": sum(by_type.values()),
        "by_type": by_type,
        "by_status": by_status,
        "by_user": by_user,
    }


def _filters_from_query(
    q: Optional[str],
    users: Optional[str],
    types: Optional[str],
    statuses: Optional[str],
    date_from: Optional[datetime],
    date_to: Optional[datetime],
    pulled: Optional[bool],
) -> OperationFilters:
    split = lambda value: [v.strip() for v in (value or "").split(",") if v.strip()]  # noqa: E731
    try:
        return OperationFilters(
            q=q,
            users=split(users),
            types=[t.upper() for t in split(types)],
            statuses=[s.upper() for s in split(statuses)],
            date_from=date_from,
            date_to=date_to,
            pulled=pulled,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/operations")
async def search_operations(
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
    q: Optional[str] = None,
    users: Annotated[Optional[str], Query(description="Comma-separated usernames")] = None,
    types: Annotated[Optional[str], Query(description="Comma-separated PUSH,PULL,UPDATE,...")] = None,
    statuses: Annotated[Optional[str], Query(description="Comma-separated statuses")] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    pulled: Optional[bool] = None,
    sort_by: str = "created_at",
    sort_order: str = "desc",
    limit: Annotated[int, Query(ge=1, le=500)] = 50,
    offset: Annotated[int, Query(ge=0)] = 0,
):
    """One page of the filtered history plus a summary of everything that matched."""
    filters = _filters_from_query(q, users, types, statuses, date_from, date_to, pulled)
    rows = (await db.execute(
        _sorted(filtered_operations(filters), sort_by, sort_order).limit(limit).offset(offset)
    )).all()
    summary = await _summary(db, filters)
    return {
        "total": summary["total"],
        "offset": offset,
        "limit": limit,
        "operations": await _responses(db, rows),
        "summary": summary,
    }


async def _local_tz(db: AsyncSession):
    """The time zone of the people reading the files: ``reports_timezone``, else UTC."""
    try:
        return (await report_service.load_report_config(db)).tz
    except report_service.ReportConfigError:
        return timezone.utc


def _xlsx_response(data: bytes, file_name: str) -> StreamingResponse:
    ascii_name = re.sub(r"[^A-Za-z0-9._-]", "_", file_name)
    return StreamingResponse(
        io.BytesIO(data),
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(file_name)}"
        },
    )


@router.post("/operations/export")
async def export_operations(
    body: ExportRequest,
    request: Request,
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """The filtered history as XLSX, with the columns (and headers) the browser asks for."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rows = (await db.execute(
        _sorted(filtered_operations(body), body.sort_by, body.sort_order).limit(EXPORT_LIMIT)
    )).all()
    operations = await _responses(db, rows)
    tz = await _local_tz(db)

    def cell(op: dict, key: str):
        if key in ("created_at", "started_at", "completed_at"):
            value = op.get(key)
            return datetime.fromisoformat(value).astimezone(tz).replace(tzinfo=None) if value else None
        if key == "pim_status":
            return (op.get("pim_delivery") or {}).get("status")
        if key == "sync_status":
            return (op.get("remote_sync") or {}).get("status")
        if key == "has_been_pulled":
            return bool(op.get("has_been_pulled")) if op.get("type") == "PUSH" else None
        return op.get(key)

    workbook = Workbook()
    workbook.properties.creator = report_service.REPORT_CREATOR
    sheet = workbook.active
    sheet.title = "RFM"
    sheet.append([column.header for column in body.columns])
    for header in sheet[1]:
        header.font = Font(bold=True)
    for op in operations:
        sheet.append([cell(op, column.key) for column in body.columns])
    for index, column in enumerate(body.columns, start=1):
        letter = get_column_letter(index)
        if column.key in ("created_at", "started_at", "completed_at"):
            for (value,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                value.number_format = "yyyy-mm-dd hh:mm:ss"
            sheet.column_dimensions[letter].width = 20
        else:
            longest = max([len(str(column.header))] + [len(str(op_cell.value or "")) for (op_cell,) in
                          sheet.iter_rows(min_row=2, min_col=index, max_col=index)])
            sheet.column_dimensions[letter].width = min(max(longest + 2, 8), 70)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    await AuditLogger.log_admin_action(
        user_id=current_user.id,
        action="history_export",
        target="operations",
        details={"rows": len(operations), "filters": body.model_dump(mode="json", exclude={"columns"})},
        ip_address=get_client_ip(request),
        user_agent=request.headers.get("user-agent"),
    )
    stamp = datetime.now(tz).strftime("%Y-%m-%d_%H%M")
    return _xlsx_response(buffer.getvalue(), f"RFM_historia_{stamp}.xlsx")


# =============================================================================
# Classification reports
# =============================================================================


async def _report_config(db: AsyncSession) -> report_service.ReportConfig:
    try:
        return await report_service.load_report_config(db)
    except report_service.ReportConfigError as exc:
        raise HTTPException(status_code=409, detail={"error": "report_config_invalid", "message": str(exc)}) from exc


@router.get("/reports/settings")
async def report_settings(
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Report settings as configured in the Admin Panel (no secrets), and the next nightly run."""
    try:
        config = await report_service.load_report_config(db)
    except report_service.ReportConfigError as exc:
        return {"valid": False, "error": str(exc)}

    now = datetime.now(config.tz)
    next_run = datetime.combine(now.date(), config.run_time, tzinfo=config.tz)
    if next_run <= now:
        next_run += timedelta(days=1)
    settings = {**config.public(), "example_path": None}
    if config.recipients:
        try:
            settings["example_path"] = "\\".join(report_service.report_location(
                config, report_service.month_start(now.date()), config.recipients[0]
            ))
        except report_service.ReportConfigError as exc:
            settings["path_error"] = str(exc)
    return {"valid": True, **settings, "next_run": next_run.isoformat() if config.auto_enabled else None}


@router.get("/reports/runs", response_model=List[ReportRunResponse])
async def report_runs(
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
    limit: Annotated[int, Query(ge=1, le=500)] = 50,
):
    """The latest report runs, newest first."""
    runs = await db.execute(select(ReportRun).order_by(desc(ReportRun.started_at), desc(ReportRun.id)).limit(limit))
    return [ReportRunResponse.model_validate(run) for run in runs.scalars()]


@router.post("/reports/download")
async def download_report(
    body: ReportDownloadRequest,
    request: Request,
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """A classification report for any users and days, straight to the browser."""
    if body.date_to < body.date_from:
        raise HTTPException(status_code=422, detail="date_to is before date_from")
    if (body.date_to - body.date_from).days > 366:
        raise HTTPException(status_code=422, detail="A report covers at most one year")

    config = await _report_config(db)
    columns = list(config.columns)
    if body.include_user_column:
        columns.insert(0, {"header": body.user_column_header or "User", "field": "username"})
    data, rows, warning = await report_service.generate_report(
        db, config, body.usernames, body.date_from, body.date_to + timedelta(days=1), columns
    )

    # A configured person's whole month gets the name it has on the share
    first, last = body.date_from, body.date_to
    recipient = next((r for r in config.recipients if [r["username"]] == body.usernames), None)
    whole_month = first.day == 1 and last + timedelta(days=1) == report_service.next_month(first)
    if recipient and whole_month:
        _, file_name = report_service.report_location(config, first, recipient)
    else:
        names = "_".join(body.usernames) if len(body.usernames) <= 3 else f"{len(body.usernames)}_users"
        file_name = f"Raport_klasyfikacji_{names}_{first.isoformat()}_{last.isoformat()}.xlsx"

    await AuditLogger.log_admin_action(
        user_id=current_user.id,
        action="report_download",
        target="reports",
        details={"usernames": body.usernames, "from": first.isoformat(), "to": last.isoformat(), "rows": rows},
        ip_address=get_client_ip(request),
        user_agent=request.headers.get("user-agent"),
    )
    response = _xlsx_response(data, file_name)
    response.headers["X-Report-Rows"] = str(rows)
    if warning:
        response.headers["X-Report-Warning"] = quote(warning)
    return response


@router.post("/reports/run", response_model=List[ReportRunResponse])
async def run_reports_now(
    body: ReportRunRequest,
    request: Request,
    current_user: Annotated[User, Depends(require_manager)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Write the configured monthly reports for ``month`` to the share now (in the background)."""
    config = await _report_config(db)
    recipients = [r for r in config.recipients if not body.usernames or r["username"] in body.usernames]
    if not recipients:
        raise HTTPException(status_code=422, detail="No configured report recipient matches")
    year, month = (int(part) for part in body.month.split("-"))
    first = date(year, month, 1)
    if first > report_service.month_start(datetime.now(config.tz).date()):
        raise HTTPException(status_code=422, detail="That month has not started yet")
    try:
        report_service.report_location(config, first, recipients[0])
    except report_service.ReportConfigError as exc:
        raise HTTPException(status_code=409, detail={"error": "report_config_invalid", "message": str(exc)}) from exc

    ids = await report_service.start_manual_runs(recipients, first, current_user.username)
    await AuditLogger.log_admin_action(
        user_id=current_user.id,
        action="report_run",
        target="reports",
        details={"month": body.month, "recipients": [r["username"] for r in recipients], "run_ids": ids},
        ip_address=get_client_ip(request),
        user_agent=request.headers.get("user-agent"),
    )
    runs = await db.execute(select(ReportRun).where(ReportRun.id.in_(ids)).order_by(ReportRun.id))
    return [ReportRunResponse.model_validate(run) for run in runs.scalars()]
